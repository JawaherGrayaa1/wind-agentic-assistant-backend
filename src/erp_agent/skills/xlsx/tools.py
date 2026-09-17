from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Any
from erp_agent.skills.loader import SkillTool

def register_tools(db: Any = None, skill_dir: Path | None = None) -> list[SkillTool]:
    """Register Excel / Spreadsheet manipulation tools."""

    def read_xlsx(file_path: str, sheet_name: str | None = None, **kwargs: Any) -> dict[str, Any]:
        p = Path(file_path)
        if not p.exists():
            return {"found": False, "error": f"File not found: {file_path}"}
        try:
            import openpyxl
            wb = openpyxl.load_workbook(p, data_only=True)
            sheet = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
            rows = list(sheet.iter_rows(values_only=True))
            return {"found": True, "rows": rows, "row_count": len(rows), "sheets": wb.sheetnames}
        except Exception as e:
            return {"found": False, "error": str(e)}

    def export_to_excel(data: list[dict[str, Any]], output_path: str = "data/exports/export.xlsx", **kwargs: Any) -> dict[str, Any]:
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Export"
            if data:
                headers = list(data[0].keys())
                ws.append(headers)
                for item in data:
                    ws.append([item.get(h) for h in headers])
            Path(output_path).parent.mkdir(parents=True, exist_ok=True)
            wb.save(output_path)
            return {"found": True, "file_path": output_path, "row_count": len(data)}
        except Exception as e:
            return {"found": False, "error": str(e)}

    return [
        SkillTool(
            name="read_xlsx",
            description="Reads rows and sheets from an Excel (.xlsx) file.",
            parameters={"file_path": "string", "sheet_name": "string"},
            handler=read_xlsx,
        ),
        SkillTool(
            name="export_to_excel",
            description="Exports a list of structured records to an Excel file.",
            parameters={"data": "array", "output_path": "string"},
            handler=export_to_excel,
        ),
    ]
